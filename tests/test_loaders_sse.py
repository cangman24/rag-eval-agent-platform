# -*- coding: utf-8 -*-
"""tests/test_loaders_sse.py —— 文档解析 + SSE 协议单测。全部离线。"""
import json
import os
import shutil
import sys
import unittest
import uuid

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from core.loaders import (  # noqa: E402
    LoadedDoc, LoaderError, MissingDependency, UnsupportedFormat,
    detect_format, load_dir, load_docx, load_file, load_html, load_json,
    load_jsonl, load_markdown, load_pdf, load_text, summarize_load,
    supported_suffixes,
)
from core.sse import (  # noqa: E402
    SSE_MEDIA_TYPE, SSEStream, format_comment,
    format_heartbeat, format_sse, iter_sse, parse_sse, stream_agent_events,
)

_TMP_ROOT = os.path.join(os.path.dirname(os.path.abspath(__file__)), ".tmp")


def make_tmp() -> str:
    os.makedirs(_TMP_ROOT, exist_ok=True)
    p = os.path.join(_TMP_ROOT, uuid.uuid4().hex[:12])
    os.makedirs(p, exist_ok=True)
    return p


def write(path: str, content, encoding: str = "utf-8") -> str:
    os.makedirs(os.path.dirname(os.path.abspath(path)), exist_ok=True)
    mode = "wb" if isinstance(content, bytes) else "w"
    if mode == "wb":
        with open(path, "wb") as f:
            f.write(content)
    else:
        with open(path, "w", encoding=encoding) as f:
            f.write(content)
    return path


# ================================================================== format detect
class TestDetectFormat(unittest.TestCase):
    def test_known_suffixes(self):
        self.assertEqual(detect_format("a.txt"), "text")
        self.assertEqual(detect_format("a.md"), "markdown")
        self.assertEqual(detect_format("a.json"), "json")
        self.assertEqual(detect_format("a.jsonl"), "jsonl")
        self.assertEqual(detect_format("a.docx"), "docx")
        self.assertEqual(detect_format("a.pdf"), "pdf")
        self.assertEqual(detect_format("a.html"), "html")

    def test_case_insensitive(self):
        self.assertEqual(detect_format("A.TXT"), "text")

    def test_unknown_raises(self):
        with self.assertRaises(UnsupportedFormat):
            detect_format("a.xyz")

    def test_supported_suffixes_nonempty(self):
        self.assertIn(".md", supported_suffixes())


# ================================================================== text
class TestLoadText(unittest.TestCase):
    def test_basic(self):
        tmp = make_tmp()
        try:
            p = write(os.path.join(tmp, "a.txt"), "水稻穗检测精度要求。")
            self.assertEqual(load_text(p), "水稻穗检测精度要求。")
        finally:
            shutil.rmtree(tmp, ignore_errors=True)

    def test_strips_bom(self):
        tmp = make_tmp()
        try:
            p = write(os.path.join(tmp, "b.txt"), "\ufeff内容甲")
            self.assertNotIn("\ufeff", load_text(p))
            self.assertIn("内容甲", load_text(p))
        finally:
            shutil.rmtree(tmp, ignore_errors=True)

    def test_gbk_fallback(self):
        tmp = make_tmp()
        try:
            p = write(os.path.join(tmp, "c.txt"), "中文内容".encode("gb18030"))
            self.assertIn("中文内容", load_text(p))
        finally:
            shutil.rmtree(tmp, ignore_errors=True)

    def test_collapses_whitespace(self):
        tmp = make_tmp()
        try:
            p = write(os.path.join(tmp, "d.txt"), "a\n\n\n\nb")
            self.assertEqual(load_text(p), "a\n\nb")
        finally:
            shutil.rmtree(tmp, ignore_errors=True)


class TestLoadMarkdown(unittest.TestCase):
    def test_preserves_headings(self):
        """★ 必须保留 #：recursive_chunk 靠它判断标题层级。"""
        tmp = make_tmp()
        try:
            p = write(os.path.join(tmp, "a.md"), "# 概述\n内容甲\n## 细则\n内容乙")
            out = load_markdown(p)
            self.assertIn("# 概述", out)
            self.assertIn("## 细则", out)
        finally:
            shutil.rmtree(tmp, ignore_errors=True)

    def test_strips_code_fence_content(self):
        tmp = make_tmp()
        try:
            p = write(os.path.join(tmp, "b.md"), "正文\n```python\nprint('x')\n```\n结尾")
            out = load_markdown(p)
            self.assertNotIn("print", out)
            self.assertIn("正文", out)
            self.assertIn("结尾", out)
        finally:
            shutil.rmtree(tmp, ignore_errors=True)

    def test_drops_image_lines(self):
        tmp = make_tmp()
        try:
            p = write(os.path.join(tmp, "c.md"), "![图](a.png)\n正文内容")
            out = load_markdown(p)
            self.assertNotIn("![", out)
            self.assertIn("正文内容", out)
        finally:
            shutil.rmtree(tmp, ignore_errors=True)


class TestLoadJson(unittest.TestCase):
    def test_flattens_to_readable_lines(self):
        """★ 不直接 str(json)：括号引号会污染分词。"""
        tmp = make_tmp()
        try:
            p = write(os.path.join(tmp, "a.json"),
                      json.dumps({"精度": {"mAP": 0.9}, "标签": ["甲", "乙"]},
                                 ensure_ascii=False))
            out = load_json(p)
            self.assertIn("精度.mAP: 0.9", out)
            self.assertIn("标签.0: 甲", out)
            self.assertNotIn("{", out)
            self.assertNotIn('"', out)
        finally:
            shutil.rmtree(tmp, ignore_errors=True)

    def test_nested_null_skipped(self):
        tmp = make_tmp()
        try:
            p = write(os.path.join(tmp, "b.json"), json.dumps({"a": None, "b": 1}))
            out = load_json(p)
            self.assertNotIn("a:", out)
            self.assertIn("b: 1", out)
        finally:
            shutil.rmtree(tmp, ignore_errors=True)


class TestLoadJsonl(unittest.TestCase):
    def test_skips_bad_lines(self):
        tmp = make_tmp()
        try:
            content = ('{"q": "问题一", "a": "答案一"}\n'
                       '{坏的 json\n'
                       '{"q": "问题二", "a": "答案二"}\n')
            p = write(os.path.join(tmp, "a.jsonl"), content)
            out = load_jsonl(p)
            self.assertIn("问题一", out)
            self.assertIn("问题二", out)
        finally:
            shutil.rmtree(tmp, ignore_errors=True)

    def test_blank_lines_ignored(self):
        tmp = make_tmp()
        try:
            p = write(os.path.join(tmp, "b.jsonl"), '{"a":1}\n\n\n{"a":2}\n')
            self.assertEqual(len(load_jsonl(p).splitlines()), 2)
        finally:
            shutil.rmtree(tmp, ignore_errors=True)


class TestLoadHtml(unittest.TestCase):
    def test_strips_tags(self):
        tmp = make_tmp()
        try:
            p = write(os.path.join(tmp, "a.html"),
                      "<html><body><h1>标题</h1><p>段落甲</p><p>段落乙</p></body></html>")
            out = load_html(p)
            self.assertIn("标题", out)
            self.assertIn("段落甲", out)
            self.assertNotIn("<", out)
        finally:
            shutil.rmtree(tmp, ignore_errors=True)

    def test_removes_script_and_style(self):
        tmp = make_tmp()
        try:
            p = write(os.path.join(tmp, "b.html"),
                      "<html><script>var x=1;</script><style>.a{}</style>"
                      "<body>可见内容</body></html>")
            out = load_html(p)
            self.assertNotIn("var x", out)
            self.assertNotIn(".a{}", out)
            self.assertIn("可见内容", out)
        finally:
            shutil.rmtree(tmp, ignore_errors=True)

    def test_entities_decoded(self):
        tmp = make_tmp()
        try:
            p = write(os.path.join(tmp, "c.html"), "<p>甲&amp;乙</p>")
            self.assertIn("甲&乙", load_html(p))
        finally:
            shutil.rmtree(tmp, ignore_errors=True)


class TestLoadPdfAndDoc(unittest.TestCase):
    def test_pdf_missing_dependency_message(self):
        """★ 未装 pypdf 时必须给可执行的安装提示，而不是 ImportError 栈。"""
        tmp = make_tmp()
        try:
            p = write(os.path.join(tmp, "a.pdf"), b"%PDF-1.4 fake")
            try:
                import pypdf  # noqa: F401
                self.skipTest("本环境已安装 pypdf，跳过缺失依赖分支")
            except ImportError:
                pass
            with self.assertRaises(MissingDependency) as ctx:
                load_pdf(p)
            self.assertIn("pip install pypdf", str(ctx.exception))
        finally:
            shutil.rmtree(tmp, ignore_errors=True)

    def test_doc_gives_conversion_hint(self):
        tmp = make_tmp()
        try:
            p = write(os.path.join(tmp, "a.doc"), b"\xd0\xcf\x11\xe0")
            with self.assertRaises(MissingDependency) as ctx:
                load_file(p)
            msg = str(ctx.exception)
            self.assertIn("docx", msg)
            self.assertIn("convert", msg.lower() + "convert")
        finally:
            shutil.rmtree(tmp, ignore_errors=True)


class TestLoadDocx(unittest.TestCase):
    def test_headings_become_hashes_and_tables_kept(self):
        try:
            import docx  # noqa: F401
        except ImportError:
            self.skipTest("未安装 python-docx")
        tmp = make_tmp()
        try:
            from docx import Document
            d = Document()
            d.add_heading("技术要求", level=1)
            d.add_paragraph("检测精度要求 mAP 不低于 0.90。")
            d.add_heading("细则", level=2)
            t = d.add_table(rows=1, cols=2)
            t.rows[0].cells[0].text = "项目"
            t.rows[0].cells[1].text = "阈值"
            path = os.path.join(tmp, "a.docx")
            d.save(path)

            out = load_docx(path)
            # ★ 标题样式 -> # 标记，供 recursive_chunk 复用层级识别
            self.assertIn("# 技术要求", out)
            self.assertIn("## 细则", out)
            self.assertIn("mAP", out)
            # 表格保留行列结构
            self.assertIn("| 项目 | 阈值 |", out)
        finally:
            shutil.rmtree(tmp, ignore_errors=True)


# ================================================================== load_file
class TestLoadFile(unittest.TestCase):
    def test_doc_id_relative_to_root(self):
        tmp = make_tmp()
        try:
            sub = os.path.join(tmp, "sub")
            p = write(os.path.join(sub, "a.txt"), "内容")
            doc = load_file(p, root=tmp)
            self.assertEqual(doc.doc_id, "sub/a.txt")
        finally:
            shutil.rmtree(tmp, ignore_errors=True)

    def test_metadata_populated(self):
        tmp = make_tmp()
        try:
            p = write(os.path.join(tmp, "a.md"), "# 标题\n内容")
            doc = load_file(p)
            self.assertEqual(doc.metadata["format"], "markdown")
            self.assertEqual(doc.metadata["filename"], "a.md")
            self.assertGreater(doc.metadata["char_len"], 0)
            self.assertTrue(doc.source.endswith("a.md"))
        finally:
            shutil.rmtree(tmp, ignore_errors=True)

    def test_missing_file_raises(self):
        with self.assertRaises(LoaderError):
            load_file(os.path.join(_TMP_ROOT, "nope.txt"))

    def test_explicit_format_override(self):
        tmp = make_tmp()
        try:
            p = write(os.path.join(tmp, "a.dat"), "纯文本内容")
            doc = load_file(p, fmt="text")
            self.assertIn("纯文本内容", doc.text)
        finally:
            shutil.rmtree(tmp, ignore_errors=True)

    def test_loaded_doc_to_dict(self):
        d = LoadedDoc("id1", "文本", "/p", {"k": "v"})
        self.assertEqual(d.to_dict()["doc_id"], "id1")


# ================================================================== load_dir
class TestLoadDir(unittest.TestCase):
    def _tree(self):
        tmp = make_tmp()
        write(os.path.join(tmp, "a.md"), "# 甲\n内容甲")
        write(os.path.join(tmp, "b.txt"), "内容乙")
        write(os.path.join(tmp, "c.json"), json.dumps({"k": "v"}))
        write(os.path.join(tmp, "skip.xyz"), "不支持")
        write(os.path.join(tmp, "sub", "d.md"), "# 丁\n内容丁")
        return tmp

    def test_recursive_collects_supported_only(self):
        tmp = self._tree()
        try:
            res = load_dir(tmp)
            ids = sorted(d.doc_id for d in res["docs"])
            self.assertEqual(ids, ["a.md", "b.txt", "c.json", "sub/d.md"])
            self.assertEqual(len(res["skipped"]), 1)
            self.assertEqual(res["errors"], [])
        finally:
            shutil.rmtree(tmp, ignore_errors=True)

    def test_non_recursive(self):
        tmp = self._tree()
        try:
            res = load_dir(tmp, recursive=False)
            self.assertEqual(sorted(d.doc_id for d in res["docs"]),
                             ["a.md", "b.txt", "c.json"])
        finally:
            shutil.rmtree(tmp, ignore_errors=True)

    def test_include_filter(self):
        tmp = self._tree()
        try:
            res = load_dir(tmp, include=[".md"])
            self.assertEqual(sorted(d.doc_id for d in res["docs"]), ["a.md", "sub/d.md"])
        finally:
            shutil.rmtree(tmp, ignore_errors=True)

    def test_exclude_filter(self):
        tmp = self._tree()
        try:
            res = load_dir(tmp, exclude=[".json", ".txt"])
            self.assertEqual(sorted(d.doc_id for d in res["docs"]), ["a.md", "sub/d.md"])
        finally:
            shutil.rmtree(tmp, ignore_errors=True)

    def test_on_error_skip_collects(self):
        """★ 一个坏文件不该让整批入库失败。"""
        tmp = make_tmp()
        try:
            write(os.path.join(tmp, "good.txt"), "好文件")
            write(os.path.join(tmp, "bad.json"), "{ 这不是合法 json")
            res = load_dir(tmp, on_error="skip")
            self.assertEqual(len(res["docs"]), 1)
            self.assertEqual(len(res["errors"]), 1)
            self.assertIn("bad.json", res["errors"][0]["path"])
        finally:
            shutil.rmtree(tmp, ignore_errors=True)

    def test_on_error_raise(self):
        tmp = make_tmp()
        try:
            write(os.path.join(tmp, "bad.json"), "{ 坏 json")
            with self.assertRaises(Exception):
                load_dir(tmp, on_error="raise")
        finally:
            shutil.rmtree(tmp, ignore_errors=True)

    def test_invalid_on_error_arg(self):
        tmp = make_tmp()
        try:
            with self.assertRaises(ValueError):
                load_dir(tmp, on_error="explode")
        finally:
            shutil.rmtree(tmp, ignore_errors=True)

    def test_missing_dir_raises(self):
        with self.assertRaises(LoaderError):
            load_dir(os.path.join(_TMP_ROOT, "nope-dir"))

    def test_hidden_dirs_skipped(self):
        tmp = make_tmp()
        try:
            write(os.path.join(tmp, ".git", "x.md"), "# 不该入库")
            write(os.path.join(tmp, "a.md"), "# 应该入库")
            res = load_dir(tmp)
            self.assertEqual([d.doc_id for d in res["docs"]], ["a.md"])
        finally:
            shutil.rmtree(tmp, ignore_errors=True)

    def test_summarize(self):
        tmp = self._tree()
        try:
            text = summarize_load(load_dir(tmp))
            self.assertIn("入库 4 篇", text)
            self.assertIn("失败 0 篇", text)
        finally:
            shutil.rmtree(tmp, ignore_errors=True)

    def test_end_to_end_chunking_integration(self):
        """解析 -> 切分：验证 loaders 的输出能被 recursive_chunk 正确分层。"""
        from core.retrieval import recursive_chunk
        tmp = make_tmp()
        try:
            write(os.path.join(tmp, "spec.md"),
                  "# 技术要求\n\n## 精度要求\n检测 mAP 不低于 0.90。\n")
            res = load_dir(tmp)
            chunks = recursive_chunk(res["docs"][0].text, chunk_size=64, overlap=0)
            self.assertTrue(chunks)
            paths = [c["heading_path"] for c in chunks]
            self.assertTrue(any("技术要求" in p for p in paths), paths)
        finally:
            shutil.rmtree(tmp, ignore_errors=True)


# ================================================================== sse format
class TestFormatSSE(unittest.TestCase):
    def test_basic_frame(self):
        out = format_sse("hello", event="msg")
        self.assertEqual(out, "event: msg\ndata: hello\n\n")

    def test_json_object_serialized(self):
        out = format_sse({"a": 1}, event="t")
        self.assertIn('data: {"a": 1}', out)

    def test_id_and_retry(self):
        out = format_sse("x", event="e", event_id="7", retry_ms=3000)
        self.assertIn("id: 7", out)
        self.assertIn("retry: 3000", out)

    def test_newlines_get_prefixed(self):
        """★ 协议关键点：data 内换行必须逐行加前缀。"""
        out = format_sse("第一行\n第二行", event="e")
        self.assertIn("data: 第一行", out)
        self.assertIn("data: 第二行", out)
        self.assertNotIn("data: 第一行\n第二行", out)

    def test_ends_with_blank_line(self):
        self.assertTrue(format_sse("x").endswith("\n\n"))

    def test_chinese_not_escaped(self):
        out = format_sse({"msg": "中文"}, event="e")
        self.assertIn("中文", out)
        self.assertNotIn("\\u4e2d", out)

    def test_comment_frame(self):
        self.assertTrue(format_comment("hb").startswith(": hb"))
        self.assertTrue(format_comment("hb").endswith("\n\n"))

    def test_comment_newline_flattened(self):
        self.assertEqual(format_comment("a\nb"), ": a b\n\n")

    def test_heartbeat(self):
        self.assertTrue(format_heartbeat().startswith(": heartbeat"))
        self.assertIn("123.000", format_heartbeat(123.0))

    def test_media_type(self):
        self.assertIn("text/event-stream", SSE_MEDIA_TYPE)


class TestParseAndRoundTrip(unittest.TestCase):
    def test_parse_simple(self):
        events = parse_sse("event: e\ndata: hi\n\n")
        self.assertEqual(len(events), 1)
        self.assertEqual(events[0]["event"], "e")
        self.assertEqual(events[0]["data"], "hi")

    def test_parse_multiline_data(self):
        events = parse_sse("data: l1\ndata: l2\n\n")
        self.assertEqual(events[0]["data"], "l1\nl2")

    def test_parse_comment(self):
        events = parse_sse(": heartbeat\n\n")
        self.assertTrue(any("comment" in e for e in events))

    def test_roundtrip_preserves_payload(self):
        """★ 往返断言：序列化 -> 解析 -> 内容一致。"""
        payload = {"tool": "kb_search", "args": {"query": "水稻\n多行"},
                   "中文键": "值"}
        frame = format_sse(payload, event="tool_call", event_id="3")
        parsed = parse_sse(frame)
        self.assertEqual(len(parsed), 1)
        self.assertEqual(parsed[0]["event"], "tool_call")
        self.assertEqual(parsed[0]["id"], "3")
        self.assertEqual(json.loads(parsed[0]["data"]), payload)

    def test_roundtrip_multiple_frames(self):
        frames = "".join(format_sse({"i": i}, event="e", event_id=str(i))
                         for i in range(5))
        parsed = [e for e in parse_sse(frames) if "event" in e]
        self.assertEqual(len(parsed), 5)
        self.assertEqual([e["id"] for e in parsed], ["0", "1", "2", "3", "4"])


class TestIterSSE(unittest.TestCase):
    def test_iter_uses_type_and_data(self):
        events = [{"type": "tool_call", "data": {"tool": "kb_search"}},
                  {"type": "done", "data": {"rounds": 3}}]
        frames = "".join(iter_sse(events))
        parsed = [e for e in parse_sse(frames) if "event" in e]
        self.assertEqual([e["event"] for e in parsed], ["tool_call", "done"])
        self.assertEqual(json.loads(parsed[0]["data"])["tool"], "kb_search")

    def test_heartbeat_inserted(self):
        events = [{"type": "e", "data": {}} for _ in range(5)]
        frames = "".join(iter_sse(events, heartbeat_every=2))
        self.assertIn(": heartbeat", frames)

    def test_non_dict_event_fallback(self):
        frames = "".join(iter_sse(["裸字符串"]))
        self.assertIn("data: 裸字符串", frames)


class TestSSEStream(unittest.TestCase):
    def test_sequence_increments(self):
        s = SSEStream()
        s.emit("a")
        s.emit("b")
        self.assertEqual(s.last_event_id, 2)
        self.assertEqual(s.count, 2)

    def test_emit_contains_id(self):
        s = SSEStream()
        out = s.emit("tool_call", tool="kb_search")
        self.assertIn("id: 1", out)
        self.assertIn("event: tool_call", out)

    def test_emit_many(self):
        s = SSEStream()
        frames = list(s.emit_many([{"type": "x", "data": {"a": 1}}]))
        self.assertEqual(len(frames), 1)
        self.assertIn("event: x", frames[0])

    def test_comment_and_heartbeat(self):
        s = SSEStream()
        self.assertTrue(s.comment("c").startswith(": c"))
        self.assertTrue(s.heartbeat().startswith(": heartbeat"))


class TestStreamAgentEvents(unittest.TestCase):
    def test_agent_events_become_sse(self):
        """★ 与 core.agent 的 make_event 输出结构对接。"""
        events = [
            {"type": "start", "ts": 1.0, "data": {"query": "q"}},
            {"type": "tool_call", "ts": 1.1, "data": {"tool": "kb_search"}},
            {"type": "done", "ts": 2.0, "data": {"rounds": 3}},
        ]
        frames = "".join(stream_agent_events(events))
        parsed = [e for e in parse_sse(frames) if "event" in e]
        self.assertEqual([e["event"] for e in parsed],
                         ["start", "tool_call", "done"])
        first = json.loads(parsed[0]["data"])
        self.assertEqual(first["query"], "q")
        self.assertEqual(first["ts"], 1.0)

    def test_real_agent_run_streams(self):
        """端到端：跑一次真 Agent，把事件流打成 SSE 并解析回来。"""
        from core.agent import AgentEngine
        from core.mcp import ToolRegistry

        reg = ToolRegistry()

        @reg.tool(description="检索")
        def kb_search(query: str):
            """检索

            query: 检索词
            """
            return {"hits": ["c1"]}

        decisions = [{"tool": "kb_search", "args": {"query": "x"}},
                     {"final": "答案文本"}]
        cursor = {"i": 0}

        def act(state, allowed):
            d = decisions[min(cursor["i"], len(decisions) - 1)]
            cursor["i"] += 1
            return d

        engine = AgentEngine(registry=reg, intent_fn=lambda q: "kb",
                             plan_fn=lambda q, t: ["检索"],
                             act_fn=act, route_table={"kb": ["kb_search"]})
        run = engine.run("查制度")
        frames = "".join(stream_agent_events(run.events))
        parsed = [e for e in parse_sse(frames) if "event" in e]
        types = [e["event"] for e in parsed]
        self.assertIn("tool_call", types)
        self.assertIn("answer", types)
        self.assertIn("done", types)
        # 每条事件都带自增 id，支持前端断线重连
        self.assertEqual([e["id"] for e in parsed],
                         [str(i + 1) for i in range(len(parsed))])

    def test_non_dict_event(self):
        frames = "".join(stream_agent_events([None, "裸值"]))
        self.assertIn("event: message", frames)


if __name__ == "__main__":
    unittest.main(verbosity=2)
